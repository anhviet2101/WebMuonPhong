"""Fill the user's original Word layouts without rebuilding their typography."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Row
from docx.text.paragraph import Paragraph


TEMPLATE_DIR = Path(__file__).with_name("document_templates")


def _normalize_sample_margins(document):
    """The supplied A sample contains decimal twips, which WordprocessingML forbids."""
    for section in document.sections:
        margins = section._sectPr.pgMar
        if margins is None:
            continue
        for name in ("top", "right", "bottom", "left", "header", "footer", "gutter"):
            key = qn(f"w:{name}")
            value = margins.get(key)
            if value and "." in value:
                margins.set(key, str(round(float(value))))


def _set_paragraph(paragraph, value):
    value = str(value or "")
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _save(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_mau_a(fields):
    doc = Document(TEMPLATE_DIR / "mau_a.docx")
    _normalize_sample_margins(doc)
    body = list(doc.paragraphs)
    header = doc.tables[0]
    signatures = doc.tables[1]

    _set_paragraph(header.cell(0, 0).paragraphs[1], fields.get("clubName", ""))
    _set_paragraph(header.cell(0, 1).paragraphs[-1], fields.get("issueDate", ""))
    intro = body[4]
    if intro.runs:
        intro.runs[0].text = str(fields.get("intro", ""))
        if len(intro.runs) > 1:
            intro.runs[1].text = " Kính đề nghị Quý phòng xem xét, hỗ trợ cụ thể như sau:"
    slots = fields.get("slots") or []
    time_para, location_para = body[6], body[7]
    last = location_para._p
    for index, slot in enumerate(slots):
        if index:
            time_xml = deepcopy(time_para._p)
            location_xml = deepcopy(location_para._p)
            last.addnext(time_xml)
            time_xml.addnext(location_xml)
            last = location_xml
            time_target = Paragraph(time_xml, time_para._parent)
            location_target = Paragraph(location_xml, location_para._parent)
        else:
            time_target, location_target = time_para, location_para
        _set_paragraph(time_target, f"Thời gian{f' {index + 1}' if len(slots) > 1 else ''}: {slot.get('time', '')}")
        _set_paragraph(location_target, f"Địa điểm{f' {index + 1}' if len(slots) > 1 else ''}: {slot.get('location', '')}")
    if not slots:
        _set_paragraph(time_para, "Thời gian: ")
        _set_paragraph(location_para, "Địa điểm: ")
    body[8]._p.getparent().remove(body[8]._p)  # No equipment request; commitment remains.
    _set_paragraph(body[9], f"Số lượng người tham gia: {fields.get('participants', '')}")
    signer_cell = signatures.cell(0, 2)
    _set_paragraph(signer_cell.paragraphs[1], fields.get("signerTitle", "CHỦ NHIỆM"))
    _set_paragraph(signer_cell.paragraphs[-1], fields.get("signerName", ""))
    return _save(doc)


def render_mau_b(rows, template, issue_date):
    doc = Document(TEMPLATE_DIR / "mau_b.docx")
    _normalize_sample_margins(doc)
    body = list(doc.paragraphs)
    header, schedule, signatures = doc.tables
    left_lines = str(template.get("leftHeader", "")).splitlines()
    left_cell = header.cell(0, 0)
    for index, paragraph in enumerate(left_cell.paragraphs[:3]):
        _set_paragraph(paragraph, left_lines[index] if index < len(left_lines) else "")
    right_cell = header.cell(0, 1)
    _set_paragraph(right_cell.paragraphs[0], template.get("rightHeader", ""))
    _set_paragraph(right_cell.paragraphs[-1], issue_date)
    _set_paragraph(body[2], template.get("title", "ĐƠN ĐỀ NGHỊ"))
    _set_paragraph(body[4], template.get("recipient", ""))
    _set_paragraph(body[6], template.get("intro", ""))
    _set_paragraph(body[9], template.get("commitment", ""))
    closing = str(template.get("closing", "")).splitlines()
    _set_paragraph(body[10], closing[0] if closing else "")
    _set_paragraph(body[11], "\n".join(closing[1:]))

    sample_row = deepcopy(schedule.rows[1]._tr)
    for row in list(schedule.rows)[1:]:
        schedule._tbl.remove(row._tr)
    for index, data in enumerate(rows):
        element = deepcopy(sample_row)
        schedule._tbl.append(element)
        row = _Row(element, schedule)
        for cell, value in zip(row.cells, [str(index + 1), data.get("time", ""), data.get("location", ""), data.get("organization", ""), data.get("note", "")]):
            _set_paragraph(cell.paragraphs[0], value)

    left_sign = signatures.cell(0, 0)
    left_lines = str(template.get("leftSignature", "")).splitlines()
    for index, paragraph in enumerate(left_sign.paragraphs[:2]):
        _set_paragraph(paragraph, left_lines[index] if index < len(left_lines) else "")
    right_sign = signatures.cell(0, 1)
    right_lines = str(template.get("rightSignature", "")).splitlines()
    for index, paragraph in enumerate(right_sign.paragraphs[:2]):
        _set_paragraph(paragraph, right_lines[index] if index < len(right_lines) else "")
    _set_paragraph(right_sign.paragraphs[-1], template.get("rightSignerName", ""))
    return _save(doc)
