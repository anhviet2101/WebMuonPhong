from datetime import datetime, time, timedelta
from io import BytesIO
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import ValidationError
from django.db import connection
from django.test import override_settings, SimpleTestCase
from docx import Document
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient, APITestCase

from backend.bookings.models import (
    Booking,
    BookingStatus,
    Building,
    BorrowingPolicy,
    BusinessRuleConfig,
    Campus,
    Organization,
    Notification,
    PhysicalStatus,
    Role,
    Room,
    RoomBlackout,
    UserProfile,
)
from backend.bookings.services.booking_service import (
    approve_booking,
    change_room,
    get_available_rooms,
    paper_deadline_for,
    submit_booking,
)
from backend.bookings.document_renderer import TEMPLATE_DIR, render_mau_a, render_mau_b
from backend.bookings.tasks import (
    auto_complete_past_bookings,
    auto_expire_unsubmitted_bookings,
    auto_release_expired_draft_holds,
    warn_overdue_physical_copies,
    send_booking_notification_emails,
)


User = get_user_model()


class OriginalWordTemplateTests(SimpleTestCase):
    def test_mau_a_multi_room_uses_reference_b_table_and_selected_header(self):
        slots = [
            {"bookingId": "1", "time": "Thứ Hai 18h-20h", "location": "Phòng 101"},
            {"bookingId": "2", "time": "Thứ Hai 18h-20h", "location": "Phòng 102"},
        ]
        document = Document(BytesIO(render_mau_a({"headerType": "doan", "clubName": "CLB A", "slots": slots})))
        self.assertEqual(len(document.tables), 3)
        self.assertEqual(len(document.tables[1].rows), 3)
        self.assertEqual(document.tables[1].cell(1, 2).text, "Phòng 101")
        self.assertEqual(document.tables[1].cell(2, 2).text, "Phòng 102")
        self.assertTrue(document.tables[1].cell(0, 1).paragraphs[0].runs[0].bold)
        self.assertTrue(document.tables[0].cell(0, 0).paragraphs[0].runs[0].bold)
        self.assertIn("BCH TRƯỜNG ĐẠI HỌC CÔNG NGHỆ", document.tables[0].cell(0, 0).text)
        self.assertEqual(document.tables[2].cell(0, 1).paragraphs[1].text, "ĐOÀN THANH NIÊN TRƯỜNG")

    def test_mau_a_keeps_reference_typography_and_commitment(self):
        content = render_mau_a({
            "clubName": "CLB A", "issueDate": "Hà Nội, ngày 24 tháng 9 năm 2026",
            "intro": "CLB A tổ chức sinh hoạt.", "participants": "30 người",
            "signerTitle": "CHỦ NHIỆM", "signerName": "Nguyễn Văn A",
            "slots": [{"time": "18h00 - 20h00", "location": "Phòng 101"}],
        })
        document = Document(BytesIO(content))
        original = Document(TEMPLATE_DIR / "mau_a.docx")
        self.assertEqual(document.sections[0].page_width, original.sections[0].page_width)
        self.assertEqual(document.sections[0].page_height, original.sections[0].page_height)
        self.assertEqual(document.sections[0].left_margin.twips, round(float(original.sections[0]._sectPr.pgMar.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}left"))))
        self.assertEqual(document.sections[0].right_margin, original.sections[0].right_margin)
        self.assertEqual(document.paragraphs[2].runs[0].font.size.pt, 20)
        self.assertEqual(document.paragraphs[3].runs[0].font.size.pt, 14)
        self.assertTrue(document.paragraphs[6].runs[0].bold)
        self.assertFalse(document.paragraphs[6].runs[1].bold)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("Cơ sở vật chất:", text)
        self.assertIn("hoàn trả lại thiết bị", text)
        self.assertEqual(len(document.tables), 2)

    def test_mau_b_preserves_five_columns_and_fourteen_point_header(self):
        content = render_mau_b(
            [{"time": "18h00 - 20h00", "location": "Phòng 101", "organization": "CLB A", "note": ""}],
            {"leftHeader": "ĐOÀN ĐHQGHN\nBCH TRƯỜNG ĐHCN\n***", "rightHeader": "ĐOÀN TNCS HỒ CHÍ MINH", "title": "ĐƠN ĐỀ NGHỊ", "recipient": "Kính gửi: Phòng HCQT", "intro": "Đề nghị mượn phòng.", "commitment": "Cam kết trả phòng.", "closing": "Xin cảm ơn!", "leftSignature": "Ý KIẾN\nPHÒNG HCQT", "rightSignature": "TM. BCH ĐOÀN TRƯỜNG\nUV BAN THƯỜNG VỤ", "rightSignerName": "Nguyễn Thị Hằng"},
            "Hà Nội, ngày 24 tháng 9 năm 2026",
        )
        document = Document(BytesIO(content))
        original = Document(TEMPLATE_DIR / "mau_b.docx")
        self.assertEqual(document.sections[0].page_width, original.sections[0].page_width)
        self.assertEqual(document.sections[0].page_height, original.sections[0].page_height)
        self.assertEqual(document.sections[0].top_margin, original.sections[0].top_margin)
        self.assertEqual(document.sections[0].bottom_margin, original.sections[0].bottom_margin)
        self.assertEqual(len(document.tables[1].columns), 5)
        self.assertEqual(len(document.tables[1].rows), 2)
        self.assertEqual(document.tables[1].cell(0, 0).paragraphs[0].runs[0].font.size.pt, 14)
        self.assertEqual(document.tables[1].cell(1, 3).text, "CLB A")
        self.assertEqual(document.tables[1].cell(1, 4).paragraphs[0].runs[0].font.size.pt, 14)
        self.assertNotIn("GĐ Kiều Mai", document.tables[1].cell(1, 2).text)


def future_weekday():
    today = timezone.localdate()
    day = today - timedelta(days=today.weekday()) + timedelta(days=7)
    return timezone.make_aware(datetime.combine(day, time(18, 0)))


class BookingApiTests(APITestCase):
    def _future_weekday(self, weekday=None):
        today = timezone.localdate()
        day = today - timedelta(days=today.weekday()) + timedelta(days=7 + (weekday or 0))
        return timezone.make_aware(datetime.combine(day, time(18, 0)))

    def _draft_payload(self, room=True):
        start = self._future_weekday()
        return {
            "room": self.room.id if room else None,
            "activity_name": "Draft activity",
            "description": "Draft description",
            "participant_count": 10,
            "contact_person": "Representative",
            "contact_phone": "0900000000",
            "contact_email": "club@example.com",
            "start_time": start.isoformat(),
            "end_time": (start + timedelta(hours=2)).isoformat(),
        }

    def setUp(self):
        self.client = APIClient()
        self.organization = Organization.objects.create(
            name="CLB A",
            abbreviation="A",
            type="club",
        )
        self.other_organization = Organization.objects.create(
            name="CLB B",
            abbreviation="B",
            type="club",
        )
        self.role = Role.objects.get(name="CLB_REP")
        self.user = User.objects.create_user("club-a", password="password", first_name="Club", last_name="A", email="a@example.com")
        self.other_user = User.objects.create_user("club-b", password="password", first_name="Club", last_name="B", email="b@example.com")
        UserProfile.objects.create(
            user=self.user,
            role=self.role,
            organization=self.organization,
            phone="0900000000", profile_completed_at=timezone.now(),
        )
        UserProfile.objects.create(
            user=self.other_user,
            role=self.role,
            organization=self.other_organization,
            phone="0900000001", profile_completed_at=timezone.now(),
        )

        campus = Campus.objects.create(name="Campus 1", code="C1")
        building = Building.objects.create(
            campus=campus,
            name="Building 1",
            floor_count=3,
        )
        self.room = Room.objects.create(
            building=building,
            name="101",
            floor=1,
            capacity=50,
            type="classroom",
        )

    def _booking(self, user=None, organization=None, room=None, start_offset=1):
        today = timezone.localdate()
        day = today - timedelta(days=today.weekday()) + timedelta(days=7 + start_offset - 1)
        start = timezone.make_aware(datetime.combine(day, time(18, 0)))
        return Booking.objects.create(
            organization=organization or self.organization,
            room=room or self.room,
            activity_name="Weekly meeting",
            description="Meeting",
            participant_count=10,
            contact_person="Representative",
            contact_phone="0900000000",
            contact_email="club@example.com",
            start_time=start,
            end_time=start + timedelta(hours=2),
            created_by=user or self.user,
        )

    def test_booking_list_is_limited_to_user_organization(self):
        own_booking = self._booking()
        self._booking(
            user=self.other_user,
            organization=self.other_organization,
            start_offset=2,
        )

        self.client.force_authenticate(self.user)
        response = self.client.get(reverse("booking-list"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual([item["id"] for item in response.data], [own_booking.id])

    def test_club_can_create_booking_without_supplying_organization(self):
        self.client.force_authenticate(self.user)
        start = self._future_weekday()

        response = self.client.post(
            reverse("booking-list"),
            {
                "room": self.room.id,
                "activity_name": "Club meeting",
                "description": "Meeting",
                "participant_count": 10,
                "contact_person": "Representative",
                "contact_phone": "0900000000",
                "contact_email": "club@example.com",
                "start_time": start.isoformat(),
                "end_time": (start + timedelta(hours=1)).isoformat(),
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Booking.objects.get(pk=response.data["id"]).organization, self.organization)
        self.assertEqual(response.data["status"], BookingStatus.PENDING_HOLD)

    def test_club_must_complete_profile_before_creating_booking(self):
        UserProfile.objects.filter(user=self.user).update(profile_completed_at=None)
        self.user.refresh_from_db()
        self.user.booking_profile.refresh_from_db()
        self.client.force_authenticate(self.user)
        response = self.client.post(reverse("booking-list"), self._draft_payload(), format="json")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Booking.objects.count(), 0)

    def test_club_can_only_book_next_week_and_respects_location_policy(self):
        self.client.force_authenticate(self.user)
        payload = self._draft_payload()
        earlier = timezone.localdate() + timedelta(days=1)
        if earlier.weekday() == 6:
            earlier += timedelta(days=1)
        if timezone.localdate().weekday() == 6:
            earlier += timedelta(days=7)
        payload["start_time"] = timezone.make_aware(datetime.combine(earlier, time(18))).isoformat()
        payload["end_time"] = timezone.make_aware(datetime.combine(earlier, time(20))).isoformat()
        self.assertEqual(self.client.post(reverse("booking-list"), payload, format="json").status_code, 400)
        payload = self._draft_payload()
        monday = self._future_weekday().date()
        BorrowingPolicy.objects.create(campus=self.room.building.campus, locked_weeks=[monday.isoformat()])
        self.assertEqual(self.client.post(reverse("booking-list"), payload, format="json").status_code, 400)
        BorrowingPolicy.objects.all().delete()
        BorrowingPolicy.objects.create(campus=self.room.building.campus, allowed_weekdays=[1, 2, 3, 4, 5])
        self.assertEqual(self.client.post(reverse("booking-list"), payload, format="json").status_code, 400)

    def test_club_can_book_second_week_but_not_third_week(self):
        self.client.force_authenticate(self.user)
        payload = self._draft_payload()
        start = self._future_weekday(7)
        payload.update(start_time=start.isoformat(), end_time=(start + timedelta(hours=2)).isoformat())
        self.assertEqual(self.client.post(reverse("booking-list"), payload, format="json").status_code, 201)
        BorrowingPolicy.objects.create(campus=self.room.building.campus, locked_weeks=[start.date().isoformat()])
        locked_payload = {**payload, "room": Room.objects.create(building=self.room.building, name="103", floor=1, capacity=50).id}
        self.assertEqual(self.client.post(reverse("booking-list"), locked_payload, format="json").status_code, 400)
        BorrowingPolicy.objects.all().delete()
        start = self._future_weekday(14)
        payload.update(start_time=start.isoformat(), end_time=(start + timedelta(hours=2)).isoformat())
        self.assertEqual(self.client.post(reverse("booking-list"), payload, format="json").status_code, 400)

    def test_multi_room_request_is_atomic_and_one_scan_covers_group(self):
        second = Room.objects.create(building=self.room.building, name="102", floor=1, capacity=50)
        BusinessRuleConfig.objects.create(key="max_bookings_per_week_per_org", value=1)
        self.client.force_authenticate(self.user)
        payload = {**self._draft_payload(), "room_ids": [self.room.id, second.id]}
        created = self.client.post(reverse("booking-create-batch"), payload, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(len(created.data), 2)
        self.assertEqual(created.data[0]["application_group"], created.data[1]["application_group"])
        third = Room.objects.create(building=self.room.building, name="103", floor=1, capacity=50)
        self.assertEqual(self.client.post(reverse("booking-list"), {**self._draft_payload(), "room": third.id}, format="json").status_code, 400)
        uploaded = self.client.post(reverse("booking-scan", args=[created.data[0]["id"]]), {"file": SimpleUploadedFile("signed.pdf", b"%PDF-1.4\n%%EOF", content_type="application/pdf")}, format="multipart")
        self.assertEqual(uploaded.status_code, 200, uploaded.data)
        self.assertIsNotNone(Booking.objects.get(pk=created.data[1]["id"]).scan_uploaded_at)
        self.assertEqual(self.client.get(reverse("booking-scan", args=[created.data[1]["id"]])).content, b"%PDF-1.4\n%%EOF")
        self.client.force_authenticate(self.other_user)
        self.assertEqual(self.client.get(reverse("booking-scan", args=[created.data[0]["id"]])).status_code, 404)
        admin = User.objects.create_superuser("batch-admin", "batch@example.com", "password")
        self.client.force_authenticate(admin)
        self.assertEqual(self.client.get(reverse("booking-scan", args=[created.data[0]["id"]])).status_code, 200)
        self.assertEqual(self.client.post(reverse("booking-confirm-scan", args=[created.data[0]["id"]])).status_code, 200)
        self.assertIsNotNone(Booking.objects.get(pk=created.data[1]["id"]).scan_confirmed_at)

    def test_multi_room_conflict_rolls_back_every_room(self):
        second = Room.objects.create(building=self.room.building, name="102", floor=1, capacity=50)
        existing = self._booking(room=second)
        submit_booking(existing, self.user)
        self.client.force_authenticate(self.user)
        payload = {**self._draft_payload(), "room_ids": [self.room.id, second.id]}
        rejected = self.client.post(reverse("booking-create-batch"), payload, format="json")
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(Booking.objects.count(), 1)

    def test_paper_deadline_is_previous_thursday_at_fifteen(self):
        booking_start = self._future_weekday(5)
        deadline = timezone.localtime(paper_deadline_for(booking_start))
        expected = timezone.localtime(booking_start).date() - timedelta(days=9)
        self.assertEqual(deadline.date(), expected)
        self.assertEqual((deadline.hour, deadline.minute), (15, 0))

    def test_admin_can_register_on_behalf_even_when_club_week_is_locked(self):
        monday = self._future_weekday().date()
        BorrowingPolicy.objects.create(campus=self.room.building.campus, locked_weeks=[monday.isoformat()])
        admin = User.objects.create_superuser("late-admin", "late@example.com", "password")
        self.client.force_authenticate(admin)
        payload = {**self._draft_payload(), "organization": self.other_organization.id}
        response = self.client.post(reverse("booking-list"), payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Booking.objects.get(pk=response.data["id"]).organization, self.other_organization)

    def test_scan_upload_confirmation_and_paper_receipt_are_separate(self):
        self.client.force_authenticate(self.user)
        created = self.client.post(reverse("booking-list"), self._draft_payload(), format="json")
        self.assertEqual(created.status_code, 201, created.data)
        booking_id = created.data["id"]
        uploaded = self.client.post(reverse("booking-scan", args=[booking_id]), {"file": SimpleUploadedFile("signed.pdf", b"%PDF-1.4\n%%EOF", content_type="application/pdf")}, format="multipart")
        self.assertEqual(uploaded.status_code, 200, uploaded.data)
        self.assertIsNotNone(uploaded.data["scan_uploaded_at"])
        admin = User.objects.create_superuser("scan-admin", "admin@example.com", "password")
        self.client.force_authenticate(admin)
        self.assertEqual(self.client.post(reverse("booking-approve", args=[booking_id])).status_code, 400)
        self.assertEqual(self.client.post(reverse("booking-confirm-scan", args=[booking_id])).status_code, 200)
        self.assertEqual(self.client.post(reverse("booking-confirm-physical", args=[booking_id])).status_code, 200)
        self.assertEqual(self.client.post(reverse("booking-approve", args=[booking_id])).status_code, 200)

    def test_draft_without_room_has_no_expiry_and_is_hidden_from_admin_list(self):
        self.client.force_authenticate(self.user)
        response = self.client.post(reverse("booking-create-draft"), {"activity_name": "Unfinished"}, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        booking = Booking.objects.get(pk=response.data["id"])
        self.assertEqual(booking.status, BookingStatus.DRAFT)
        self.assertIsNone(booking.room)
        self.assertIsNone(booking.hold_expires_at)
        self.assertEqual([item["id"] for item in self.client.get(reverse("booking-list")).data], [booking.id])
        admin = User.objects.create_user("draft-admin", password="password", is_staff=True)
        self.client.force_authenticate(admin)
        self.assertNotIn(booking.id, [item["id"] for item in self.client.get(reverse("booking-list")).data])

    def test_roomed_draft_holds_for_one_hour_then_releases_room(self):
        self.client.force_authenticate(self.user)
        payload = self._draft_payload()
        response = self.client.post(reverse("booking-create-draft"), payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        booking = Booking.objects.get(pk=response.data["id"])
        self.assertLess(abs((booking.hold_expires_at - timezone.now()).total_seconds() - 3600), 15)
        original_expiry = booking.hold_expires_at
        params = {"start_time": payload["start_time"], "end_time": payload["end_time"]}
        blocked = self.client.get(reverse("room-available"), params)
        self.assertNotIn(self.room.id, [item["id"] for item in blocked.data])
        edited = self.client.patch(reverse("booking-detail", args=[booking.id]), {"activity_name": "Edited draft"}, format="json")
        self.assertEqual(edited.status_code, 200, edited.data)
        booking.refresh_from_db()
        self.assertEqual(booking.hold_expires_at, original_expiry)
        booking.hold_expires_at = timezone.now() - timedelta(seconds=1)
        booking.save(update_fields=["hold_expires_at"])
        self.assertEqual(auto_release_expired_draft_holds(), 1)
        booking.refresh_from_db()
        self.assertEqual(booking.status, BookingStatus.DRAFT)
        self.assertIsNone(booking.room)
        self.assertIsNone(booking.hold_expires_at)
        self.assertIsNotNone(original_expiry)
        available = self.client.get(reverse("room-available"), params)
        self.assertIn(self.room.id, [item["id"] for item in available.data])

    def test_draft_can_release_room_without_losing_content(self):
        self.client.force_authenticate(self.user)
        response = self.client.post(reverse("booking-create-draft"), self._draft_payload(), format="json")
        self.assertEqual(response.status_code, 201, response.data)
        booking_id = response.data["id"]
        updated = self.client.patch(reverse("booking-detail", args=[booking_id]), {"room": None}, format="json")
        self.assertEqual(updated.status_code, 200, updated.data)
        booking = Booking.objects.get(pk=booking_id)
        self.assertIsNone(booking.room)
        self.assertIsNone(booking.hold_expires_at)
        self.assertEqual(booking.activity_name, "Draft activity")

    def test_submitting_roomed_draft_uses_normal_submission_deadline(self):
        self.client.force_authenticate(self.user)
        created = self.client.post(reverse("booking-create-draft"), self._draft_payload(), format="json")
        self.assertEqual(created.status_code, 201, created.data)
        submitted = self.client.patch(reverse("booking-update-and-submit", args=[created.data["id"]]), {}, format="json")
        self.assertEqual(submitted.status_code, 200, submitted.data)
        booking = Booking.objects.get(pk=created.data["id"])
        self.assertEqual(booking.status, BookingStatus.PENDING_HOLD)
        self.assertGreater((booking.scan_deadline_at - timezone.now()).total_seconds(), 23 * 3600)
        self.assertEqual(booking.hold_expires_at, booking.scan_deadline_at)

    def test_second_roomed_draft_cannot_hold_same_room_and_time(self):
        self.client.force_authenticate(self.user)
        payload = self._draft_payload()
        first = self.client.post(reverse("booking-create-draft"), payload, format="json")
        self.assertEqual(first.status_code, 201, first.data)
        second = self.client.post(reverse("booking-create-draft"), payload, format="json")
        self.assertEqual(second.status_code, 400)
        self.assertEqual(Booking.objects.filter(status=BookingStatus.DRAFT).count(), 1)

    def test_rejects_sunday_after_21_and_letters_in_phone(self):
        self.client.force_authenticate(self.user)
        sunday = timezone.localdate() + timedelta(days=(6 - timezone.localdate().weekday()) % 7 or 7)
        sunday_start = timezone.make_aware(datetime.combine(sunday, time(18, 0)))
        payload = self._draft_payload()
        payload["start_time"] = sunday_start.isoformat()
        payload["end_time"] = (sunday_start + timedelta(hours=1)).isoformat()
        self.assertEqual(self.client.post(reverse("booking-list"), payload, format="json").status_code, 400)
        weekday = self._future_weekday()
        payload["start_time"] = weekday.replace(hour=20).isoformat()
        payload["end_time"] = weekday.replace(hour=21, minute=1).isoformat()
        self.assertEqual(self.client.post(reverse("booking-list"), payload, format="json").status_code, 400)
        payload["end_time"] = weekday.replace(hour=21, minute=0).isoformat()
        payload["contact_phone"] = "0900ABC000"
        self.assertEqual(self.client.post(reverse("booking-list"), payload, format="json").status_code, 400)

    def test_create_conflict_does_not_leave_draft(self):
        held = self._booking()
        submit_booking(held, self.user)
        self.client.force_authenticate(self.user)
        response = self.client.post(reverse("booking-list"), {
            "room": self.room.id,
            "activity_name": "Conflicting meeting",
            "description": "Meeting",
            "participant_count": 10,
            "contact_person": "Representative",
            "contact_phone": "0900000000",
            "contact_email": "club@example.com",
            "start_time": held.start_time.isoformat(),
            "end_time": held.end_time.isoformat(),
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Booking.objects.filter(activity_name="Conflicting meeting").exists())

    def test_update_and_submit_draft_saves_changes_and_holds_room(self):
        booking = self._booking()
        self.client.force_authenticate(self.user)
        response = self.client.patch(
            reverse("booking-update-and-submit", args=[booking.id]),
            {"activity_name": "Updated meeting", "participant_count": 25},
            format="json",
        )
        self.assertEqual(response.status_code, 200, response.data)
        booking.refresh_from_db()
        self.assertEqual(booking.activity_name, "Updated meeting")
        self.assertEqual(booking.participant_count, 25)
        self.assertEqual(booking.status, BookingStatus.PENDING_HOLD)
        self.assertIsNotNone(booking.hold_expires_at)

    def test_update_and_submit_conflict_rolls_back_draft_edits(self):
        held = self._booking()
        submit_booking(held, self.user)
        draft = self._booking(start_offset=2)
        self.client.force_authenticate(self.user)
        response = self.client.patch(
            reverse("booking-update-and-submit", args=[draft.id]),
            {
                "activity_name": "Should not be saved",
                "start_time": held.start_time.isoformat(),
                "end_time": held.end_time.isoformat(),
            },
            format="json",
        )
        self.assertEqual(response.status_code, 400)
        draft.refresh_from_db()
        self.assertEqual(draft.status, BookingStatus.DRAFT)
        self.assertEqual(draft.activity_name, "Weekly meeting")
        self.assertNotEqual(draft.start_time, held.start_time)

    def test_admin_account_requires_and_keeps_selected_club(self):
        admin_role = Role.objects.get(name="YU_ADMIN")
        admin = User.objects.create_user("office", password="OfficePass123!")
        UserProfile.objects.create(user=admin, role=admin_role)
        self.client.force_authenticate(admin)

        missing = self.client.post(
            reverse("admin-users"),
            {"username": "new-club", "role": "CLB_REP"},
            format="json",
        )
        self.assertEqual(missing.status_code, 400)

        created = self.client.post(
            reverse("admin-users"),
            {
                "username": "new-club",
                "role": "CLB_REP",
                "organization": self.other_organization.id,
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(
            UserProfile.objects.get(user_id=created.data["id"]).organization,
            self.other_organization,
        )

    def test_approve_requires_confirmed_scan_and_hard_copy(self):
        booking = self._booking()
        submit_booking(booking, self.user)
        admin = User.objects.create_user("admin", password="password", is_staff=True)
        with self.assertRaises(ValidationError):
            approve_booking(booking, admin)
        booking.scan_uploaded_at = timezone.now()
        booking.scan_confirmed_at = timezone.now()
        booking.physical_status = PhysicalStatus.DA_NHAN_BAN_CUNG
        booking.save(update_fields=["scan_uploaded_at", "scan_confirmed_at", "physical_status"])
        approved = approve_booking(booking, admin)
        self.assertEqual(approved.status, BookingStatus.APPROVED)
        self.assertEqual(approved.physical_status, PhysicalStatus.DA_NHAN_BAN_CUNG)

    def test_available_rooms_excludes_conflicting_booking(self):
        booking = self._booking()
        submit_booking(booking, self.user)

        self.client.force_authenticate(self.user)
        response = self.client.get(
            reverse("room-available"),
            {
                "start_time": booking.start_time.isoformat(),
                "end_time": booking.end_time.isoformat(),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertNotIn(self.room.id, [item["id"] for item in response.data])

    def test_available_rooms_filters_by_capacity(self):
        self.client.force_authenticate(self.user)
        start = self._future_weekday()
        params = {
            "start_time": start.isoformat(),
            "end_time": (start + timedelta(hours=2)).isoformat(),
        }
        enough = self.client.get(reverse("room-available"), {**params, "participant_count": 50})
        too_small = self.client.get(reverse("room-available"), {**params, "participant_count": 51})
        invalid = self.client.get(reverse("room-available"), {**params, "participant_count": 0})
        self.assertEqual(enough.status_code, 200, enough.data)
        self.assertIn(self.room.id, [item["id"] for item in enough.data])
        self.assertEqual(too_small.status_code, 200, too_small.data)
        self.assertNotIn(self.room.id, [item["id"] for item in too_small.data])
        self.assertEqual(invalid.status_code, 400)

    def test_available_rooms_query_count_does_not_grow_with_rooms(self):
        Room.objects.bulk_create([
            Room(building=self.room.building, name=f"Extra {index}", floor=1, capacity=50)
            for index in range(12)
        ])
        start = self._future_weekday()
        with CaptureQueriesContext(connection) as queries:
            available = get_available_rooms(
                start, start + timedelta(hours=2),
                queryset=Room.objects.filter(building=self.room.building),
            )
        self.assertEqual(len(available), 13)
        self.assertLessEqual(len(queries), 5)

    @override_settings(EMAIL_DELIVERY_MODE="async")
    def test_submit_queues_email_only_after_commit(self):
        self.user.email = "club@example.com"
        self.user.save(update_fields=["email"])
        booking = self._booking()
        with patch("backend.bookings.tasks.send_booking_notification_emails.delay") as enqueue:
            with self.captureOnCommitCallbacks(execute=True) as callbacks:
                submit_booking(booking, self.user)
                enqueue.assert_not_called()
            self.assertEqual(len(callbacks), 1)
            enqueue.assert_called_once_with(["club@example.com"], "Đơn 'Weekly meeting' đã được gửi và đang giữ chỗ.")

    @override_settings(
        EMAIL_DELIVERY_MODE="sync",
        EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
    )
    def test_submit_sends_email_without_worker(self):
        self.user.email = "club@example.com"
        self.user.save(update_fields=["email"])
        with self.captureOnCommitCallbacks(execute=True):
            submit_booking(self._booking(), self.user)
        self.assertEqual([message.to for message in mail.outbox], [["club@example.com"]])

    @override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_email_worker_sends_individual_messages(self):
        send_booking_notification_emails(["club@example.com", "admin@example.com"], "Đơn đã được gửi")
        self.assertEqual([message.to for message in mail.outbox], [
            ["club@example.com"], ["admin@example.com"],
        ])

    def test_minimum_fifteen_minute_buffer_blocks_room(self):
        self.room.buffer_before_minutes = 0
        self.room.buffer_after_minutes = 0
        self.room.save(update_fields=["buffer_before_minutes", "buffer_after_minutes"])
        booking = self._booking()
        submit_booking(booking, self.user)
        candidate_start = booking.end_time + timedelta(minutes=20)
        candidate_end = candidate_start + timedelta(minutes=30)
        self.client.force_authenticate(self.user)
        response = self.client.get(reverse("room-available"), {
            "start_time": candidate_start.isoformat(),
            "end_time": candidate_end.isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(self.room.id, [item["id"] for item in response.data])
        candidate = Booking.objects.create(
            organization=self.organization,
            room=self.room,
            activity_name="Too close to prior booking",
            description="Test buffer",
            participant_count=10,
            contact_person="Representative",
            contact_phone="0900000000",
            contact_email="club@example.com",
            start_time=candidate_start,
            end_time=candidate_end,
            created_by=self.user,
        )
        with self.assertRaisesMessage(ValidationError, "Phòng đã có lịch trùng"):
            submit_booking(candidate, self.user)

    def test_available_rooms_allows_excluding_current_booking(self):
        booking = self._booking()
        submit_booking(booking, self.user)

        self.client.force_authenticate(self.user)
        response = self.client.get(
            reverse("room-available"),
            {
                "start_time": booking.start_time.isoformat(),
                "end_time": booking.end_time.isoformat(),
                "exclude_booking": booking.id,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(self.room.id, [item["id"] for item in response.data])

    def test_available_rooms_excludes_approved_booking(self):
        booking = self._booking()
        submit_booking(booking, self.user)
        admin = User.objects.create_user("approver", password="password", is_staff=True)
        booking.scan_uploaded_at = timezone.now()
        booking.scan_confirmed_at = timezone.now()
        booking.physical_status = PhysicalStatus.DA_NHAN_BAN_CUNG
        booking.save(update_fields=["scan_uploaded_at", "scan_confirmed_at", "physical_status"])
        approve_booking(booking, admin)
        self.client.force_authenticate(self.other_user)
        response = self.client.get(reverse("room-available"), {
            "start_time": booking.start_time.isoformat(),
            "end_time": booking.end_time.isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(self.room.id, [item["id"] for item in response.data])

    def test_available_rooms_respects_blackout_and_buffer(self):
        start = self._future_weekday()
        RoomBlackout.objects.create(
            scope_type="room",
            room_ids=[self.room.id],
            start_time=start + timedelta(hours=1, minutes=10),
            end_time=start + timedelta(hours=2),
            reason="Maintenance",
            created_by=self.user,
        )
        self.client.force_authenticate(self.user)
        response = self.client.get(reverse("room-available"), {
            "start_time": start.isoformat(),
            "end_time": (start + timedelta(hours=1)).isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(self.room.id, [item["id"] for item in response.data])

    def test_available_rooms_excludes_room_changed_booking(self):
        booking = self._booking()
        booking.status = BookingStatus.ROOM_CHANGED
        booking.save(update_fields=["status"])
        self.client.force_authenticate(self.other_user)
        response = self.client.get(reverse("room-available"), {
            "start_time": booking.start_time.isoformat(),
            "end_time": booking.end_time.isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(self.room.id, [item["id"] for item in response.data])

    def test_available_rooms_excludes_building_blackout(self):
        start = self._future_weekday()
        RoomBlackout.objects.create(
            scope_type="building",
            room_ids=[],
            building=self.room.building,
            start_time=start,
            end_time=start + timedelta(hours=1),
            reason="Exam",
            created_by=self.user,
        )
        self.client.force_authenticate(self.user)
        response = self.client.get(reverse("room-available"), {
            "start_time": start.isoformat(),
            "end_time": (start + timedelta(hours=1)).isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(self.room.id, [item["id"] for item in response.data])

    def test_edit_cannot_move_active_booking_into_blackout(self):
        booking = self._booking()
        submit_booking(booking, self.user)
        new_start = booking.start_time + timedelta(hours=4)
        RoomBlackout.objects.create(
            scope_type="room",
            room_ids=[self.room.id],
            start_time=new_start,
            end_time=new_start + timedelta(hours=1),
            reason="Maintenance",
            created_by=self.user,
        )
        self.client.force_authenticate(self.user)
        response = self.client.patch(
            reverse("booking-detail", args=[booking.id]),
            {
                "start_time": new_start.isoformat(),
                "end_time": (new_start + timedelta(hours=1)).isoformat(),
            },
            format="json",
        )
        self.assertEqual(response.status_code, 400, response.data)
        booking.refresh_from_db()
        self.assertNotEqual(booking.start_time, new_start)

    def test_calendar_redacts_other_club_details(self):
        booking = self._booking(
            user=self.other_user,
            organization=self.other_organization,
        )
        submit_booking(booking, self.other_user)
        self.client.force_authenticate(self.user)
        response = self.client.get(reverse("booking-calendar"))
        self.assertEqual(response.status_code, 200)
        item = next(item for item in response.data if item.get("hidden_details"))
        self.assertTrue(item["hidden_details"])
        self.assertNotEqual(item["id"], booking.id)
        self.assertEqual(item["activity_name"], "Không khả dụng")
        self.assertEqual(item["organization_name"], "")

    def test_same_club_representative_can_cancel_booking(self):
        booking = self._booking()
        submit_booking(booking, self.user)
        colleague = User.objects.create_user("club-a-colleague", password="password")
        UserProfile.objects.create(
            user=colleague,
            role=self.role,
            organization=self.organization,
        )
        self.client.force_authenticate(colleague)
        response = self.client.post(reverse("booking-cancel", args=[booking.id]))
        self.assertEqual(response.status_code, 200, response.data)
        booking.refresh_from_db()
        self.assertEqual(booking.status, BookingStatus.CANCELLED)

    def test_admin_can_change_pending_room_without_approving(self):
        booking = self._booking()
        submit_booking(booking, self.user)
        booking.refresh_from_db()
        other_room = Room.objects.create(
            building=self.room.building,
            name="102",
            floor=1,
            capacity=50,
            type="classroom",
        )
        admin = User.objects.create_user("office", password="password", is_staff=True)
        changed = change_room(booking, other_room, admin, reason="Room adjustment")
        self.assertEqual(changed.status, BookingStatus.PENDING_HOLD)
        self.assertEqual(changed.room_id, other_room.id)
        self.assertEqual(changed.hold_expires_at, booking.hold_expires_at)

    def test_approval_rejects_new_blackout(self):
        booking = self._booking()
        submit_booking(booking, self.user)
        admin = User.objects.create_user("office", password="password", is_staff=True)
        RoomBlackout.objects.create(
            scope_type="room",
            room_ids=[self.room.id],
            start_time=booking.start_time,
            end_time=booking.end_time,
            reason="Maintenance",
            created_by=admin,
        )
        with self.assertRaises(ValidationError):
            approve_booking(booking, admin)

    def test_submit_rejects_overlapping_room_booking(self):
        first = self._booking()
        submit_booking(first, self.user)
        second = self._booking(start_offset=1)
        second.start_time = first.start_time + timedelta(minutes=30)
        second.end_time = first.end_time + timedelta(minutes=30)
        second.save(update_fields=["start_time", "end_time", "updated_at"])

        with self.assertRaises(ValidationError):
            submit_booking(second, self.user)

    def test_club_cannot_edit_after_physical_copy_confirmed(self):
        booking = self._booking()
        booking.status = BookingStatus.NEEDS_REVISION
        booking.physical_status = PhysicalStatus.DA_NHAN_BAN_CUNG
        booking.save(update_fields=["status", "physical_status", "updated_at"])

        self.client.force_authenticate(self.user)
        response = self.client.patch(
            reverse("booking-detail", args=[booking.id]),
            {"activity_name": "Updated after hard copy"},
            format="json",
        )

        self.assertEqual(response.status_code, 403)
        booking.refresh_from_db()
        self.assertEqual(booking.activity_name, "Weekly meeting")

    def test_club_can_edit_pending_hold_before_physical_copy(self):
        booking = self._booking()
        submit_booking(booking, self.user)

        self.client.force_authenticate(self.user)
        response = self.client.patch(
            reverse("booking-detail", args=[booking.id]),
            {"activity_name": "Updated before hard copy"},
            format="json",
        )

        self.assertEqual(response.status_code, 200)
        booking.refresh_from_db()
        self.assertEqual(booking.activity_name, "Updated before hard copy")


class BookingConflictServiceTests(APITestCase):
    def setUp(self):
        self.organization = Organization.objects.create(
            name="CLB",
            abbreviation="CLB",
            type="club",
        )
        role = Role.objects.get(name="CLB_REP")
        self.user = User.objects.create_user("club", password="password")
        UserProfile.objects.create(
            user=self.user,
            role=role,
            organization=self.organization,
            phone="0900000000", profile_completed_at=timezone.now(),
        )
        campus = Campus.objects.create(name="Campus", code="C")
        building = Building.objects.create(
            campus=campus,
            name="Building",
            floor_count=2,
        )
        self.room = Room.objects.create(
            building=building,
            name="101",
            floor=1,
            capacity=20,
            type="classroom",
        )

    def test_submit_rejects_overlapping_hold(self):
        start = future_weekday()
        first = Booking.objects.create(
            organization=self.organization,
            room=self.room,
            activity_name="First",
            description="First",
            participant_count=5,
            contact_person="Person",
            contact_phone="0900000000",
            contact_email="club@example.com",
            start_time=start,
            end_time=start + timedelta(hours=1),
            created_by=self.user,
        )
        second = Booking.objects.create(
            organization=self.organization,
            room=self.room,
            activity_name="Second",
            description="Second",
            participant_count=5,
            contact_person="Person",
            contact_phone="0900000000",
            contact_email="club@example.com",
            start_time=start + timedelta(minutes=10),
            end_time=start + timedelta(hours=1, minutes=10),
            created_by=self.user,
        )

        submit_booking(first, self.user)
        with self.assertRaisesMessage(
            ValidationError,
            "Phòng đã có lịch trùng trong khoảng thời gian này.",
        ):
            submit_booking(second, self.user)

        second.refresh_from_db()
        self.assertEqual(second.status, BookingStatus.DRAFT)


class BookingMaintenanceTaskTests(APITestCase):
    def setUp(self):
        self.organization = Organization.objects.create(
            name="CLB",
            abbreviation="CLB",
            type="club",
        )
        role = Role.objects.get(name="CLB_REP")
        self.user = User.objects.create_user("club", password="password")
        UserProfile.objects.create(
            user=self.user,
            role=role,
            organization=self.organization,
        )
        campus = Campus.objects.create(name="Campus", code="C")
        building = Building.objects.create(
            campus=campus,
            name="Building",
            floor_count=2,
        )
        self.room = Room.objects.create(
            building=building,
            name="101",
            floor=1,
            capacity=20,
            type="classroom",
        )

    def test_expire_task_updates_booking_audit_and_notification(self):
        start = future_weekday()
        booking = Booking.objects.create(
            organization=self.organization,
            room=self.room,
            activity_name="Expired request",
            description="Request",
            participant_count=5,
            contact_person="Person",
            contact_phone="0900000000",
            contact_email="club@example.com",
            start_time=start,
            end_time=start + timedelta(hours=1),
            status=BookingStatus.PENDING_HOLD,
            hold_expires_at=timezone.now() - timedelta(minutes=1),
            scan_deadline_at=timezone.now() - timedelta(minutes=1),
            created_by=self.user,
        )

        self.assertEqual(auto_expire_unsubmitted_bookings(), 1)

        booking.refresh_from_db()
        self.assertEqual(booking.status, BookingStatus.EXPIRED)
        self.assertTrue(booking.created_by.audit_logs.filter(entity_id=str(booking.id)).exists())
        self.assertTrue(
            Notification.objects.filter(
                user=self.user,
                related_booking=booking,
                type=Notification.NotificationType.EXPIRED,
            ).exists()
        )

    def test_complete_task_updates_past_booking(self):
        end = timezone.now() - timedelta(minutes=1)
        booking = Booking.objects.create(
            organization=self.organization,
            room=self.room,
            activity_name="Past booking",
            description="Booking",
            participant_count=5,
            contact_person="Person",
            contact_phone="0900000000",
            contact_email="club@example.com",
            start_time=end - timedelta(hours=1),
            end_time=end,
            status=BookingStatus.APPROVED,
            created_by=self.user,
        )

        self.assertEqual(auto_complete_past_bookings(), 1)

        booking.refresh_from_db()
        self.assertEqual(booking.status, BookingStatus.COMPLETED)

    def test_overdue_paper_warns_staff_without_cancelling_room(self):
        admin = User.objects.create_user("office-warning", password="password")
        UserProfile.objects.create(user=admin, role=Role.objects.get(name="YU_ADMIN"))
        start = future_weekday()
        booking = Booking.objects.create(
            organization=self.organization, room=self.room, activity_name="Paper pending",
            description="Meeting", participant_count=5, contact_person="Person",
            contact_phone="0900000000", contact_email="club@example.com",
            start_time=start, end_time=start + timedelta(hours=1),
            status=BookingStatus.PENDING_HOLD, created_by=self.user,
            scan_uploaded_at=timezone.now(), scan_deadline_at=timezone.now() - timedelta(hours=1),
            paper_deadline_at=timezone.now() - timedelta(minutes=1),
        )
        self.assertEqual(warn_overdue_physical_copies(), 1)
        self.assertEqual(warn_overdue_physical_copies(), 0)
        self.assertEqual(auto_expire_unsubmitted_bookings(), 0)
        booking.refresh_from_db()
        self.assertEqual(booking.status, BookingStatus.PENDING_HOLD)
        self.assertTrue(Notification.objects.filter(user=admin, related_booking=booking, type=Notification.NotificationType.PHYSICAL_REMINDER).exists())
